"""Ingestion pipeline: parse QMS documents, chunk, extract cross-references, and build indexes."""

from __future__ import annotations

import json
import logging
import os
import pickle
import re
from dataclasses import dataclass, field, asdict
from pathlib import Path

import lancedb
import networkx as nx
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from openai import OpenAI

logger = logging.getLogger(__name__)

DOC_TYPE_PREFIXES = (
    "3P", "BOM", "DHF", "DMR", "DR", "ECR", "ESF", "IFU",
    "MEMO", "PLN", "QSR", "RSK", "TRA", "VVAM", "VVPR",
)

DOC_ID_PATTERN = re.compile(
    r"(?P<prefix>" + "|".join(DOC_TYPE_PREFIXES) + r")"
    r"[-\s]*(?:(?:M02|MC2|SWV)[-\s]*)?(?:\d+|MC2)",
    re.IGNORECASE,
)

REVISION_PATTERN = re.compile(r"_([A-Z])(?:[-_]|$)", re.IGNORECASE)
OBSOLETE_PATTERN = re.compile(r"[-_]Obsolete", re.IGNORECASE)
SIGNED_PATTERN = re.compile(r"[-_][Ss]igned")

FALLBACK_SPLITTER = RecursiveCharacterTextSplitter(
    chunk_size=1500,
    chunk_overlap=200,
    separators=["\n\n", "\n", ". ", " ", ""],
)


@dataclass
class DocMetadata:
    doc_id: str
    doc_type: str
    title: str
    revision: str
    status: str  # "current" | "obsolete" | "signed"
    file_format: str  # "docx" | "xlsx" | "unknown"
    filename: str


@dataclass
class Chunk:
    text: str
    doc_id: str
    doc_type: str
    title: str
    revision: str
    status: str
    section: str
    chunk_index: int
    filename: str


@dataclass
class CrossReference:
    source_doc_id: str
    target_doc_id: str
    context: str
    source_file: str


# ---------------------------------------------------------------------------
# Filename metadata extraction
# ---------------------------------------------------------------------------

def parse_filename(filename: str) -> DocMetadata:
    """Extract structured metadata from a QMS filename."""
    stem = Path(filename).stem if "." in filename else filename

    ext = Path(filename).suffix.lower() if "." in filename else ""
    if ext in (".docx",):
        file_format = "docx"
    elif ext in (".xlsx",):
        file_format = "xlsx"
    else:
        file_format = "unknown"

    raw_id = stem.split(" - ")[0].strip() if " - " in stem else stem
    doc_id = re.sub(r"\s+", "-", raw_id)

    type_match = re.match(
        r"(?P<prefix>" + "|".join(DOC_TYPE_PREFIXES) + r")",
        doc_id,
        re.IGNORECASE,
    )
    doc_type = type_match.group("prefix").upper() if type_match else "UNKNOWN"

    parts = stem.split(" - ", 1)
    title = parts[1].strip() if len(parts) > 1 else parts[0].strip()
    rev_in_title = REVISION_PATTERN.search(title)
    if rev_in_title:
        title = title[:rev_in_title.start()].strip()

    rev_match = REVISION_PATTERN.search(stem)
    revision = rev_match.group(1).upper() if rev_match else ""

    if OBSOLETE_PATTERN.search(stem):
        status = "obsolete"
    elif SIGNED_PATTERN.search(stem):
        status = "signed"
    else:
        status = "current"

    return DocMetadata(
        doc_id=doc_id,
        doc_type=doc_type,
        title=title,
        revision=revision,
        status=status,
        file_format=file_format,
        filename=filename,
    )


# ---------------------------------------------------------------------------
# Document parsing
# ---------------------------------------------------------------------------

def parse_docx(filepath: Path) -> list[tuple[str, str]]:
    """Parse a .docx file into (heading, text) sections.

    Returns a list of (section_heading, section_text) tuples.
    """
    try:
        doc = DocxDocument(str(filepath))
    except Exception:
        logger.warning("Failed to parse docx: %s", filepath.name)
        return []

    sections: list[tuple[str, str]] = []
    current_heading = ""
    current_paragraphs: list[str] = []

    for para in doc.paragraphs:
        style_name = (para.style.name or "").lower() if para.style else ""
        text = para.text.strip()
        if not text:
            continue

        if "heading" in style_name:
            if current_paragraphs:
                sections.append((current_heading, "\n".join(current_paragraphs)))
                current_paragraphs = []
            current_heading = text
        else:
            current_paragraphs.append(text)

    if current_paragraphs:
        sections.append((current_heading, "\n".join(current_paragraphs)))

    for table in doc.tables:
        rows_text = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows_text.append(" | ".join(cells))
        if rows_text:
            table_text = "\n".join(rows_text)
            sections.append(("Table", table_text))

    if not sections:
        full_text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
        if full_text:
            sections.append(("", full_text))

    return sections


def parse_xlsx(filepath: Path) -> list[tuple[str, str]]:
    """Parse an .xlsx file into (sheet_name, text) sections."""
    try:
        wb = load_workbook(str(filepath), read_only=True, data_only=True)
    except Exception:
        logger.warning("Failed to parse xlsx: %s", filepath.name)
        return []

    sections: list[tuple[str, str]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            sections.append((sheet_name, "\n".join(rows)))

    wb.close()
    return sections


def parse_document(filepath: Path) -> list[tuple[str, str]]:
    """Parse a document file, auto-detecting format."""
    ext = filepath.suffix.lower()
    if ext == ".docx":
        return parse_docx(filepath)
    elif ext == ".xlsx":
        return parse_xlsx(filepath)
    else:
        result = parse_docx(filepath)
        if result:
            return result
        return parse_xlsx(filepath)


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------

def chunk_document(
    meta: DocMetadata,
    sections: list[tuple[str, str]],
) -> list[Chunk]:
    """Split parsed document sections into chunks with metadata."""
    chunks: list[Chunk] = []
    chunk_idx = 0

    for heading, text in sections:
        if not text.strip():
            continue

        if len(text) > 2000:
            sub_chunks = FALLBACK_SPLITTER.split_text(text)
            for sc in sub_chunks:
                chunks.append(Chunk(
                    text=sc,
                    doc_id=meta.doc_id,
                    doc_type=meta.doc_type,
                    title=meta.title,
                    revision=meta.revision,
                    status=meta.status,
                    section=heading,
                    chunk_index=chunk_idx,
                    filename=meta.filename,
                ))
                chunk_idx += 1
        else:
            chunks.append(Chunk(
                text=text,
                doc_id=meta.doc_id,
                doc_type=meta.doc_type,
                title=meta.title,
                revision=meta.revision,
                status=meta.status,
                section=heading,
                chunk_index=chunk_idx,
                filename=meta.filename,
            ))
            chunk_idx += 1

    return chunks


# ---------------------------------------------------------------------------
# Cross-reference extraction
# ---------------------------------------------------------------------------

def extract_cross_references(
    meta: DocMetadata,
    full_text: str,
) -> list[CrossReference]:
    """Extract document ID references from text content."""
    refs: list[CrossReference] = []
    seen: set[str] = set()

    for match in DOC_ID_PATTERN.finditer(full_text):
        target_raw = match.group(0)
        target_id = re.sub(r"\s+", "-", target_raw.strip()).upper()

        if target_id == meta.doc_id.upper():
            continue
        if target_id in seen:
            continue
        seen.add(target_id)

        start = max(0, match.start() - 80)
        end = min(len(full_text), match.end() + 80)
        context = full_text[start:end].replace("\n", " ").strip()

        refs.append(CrossReference(
            source_doc_id=meta.doc_id,
            target_doc_id=target_id,
            context=context,
            source_file=meta.filename,
        ))

    return refs


# ---------------------------------------------------------------------------
# Embedding
# ---------------------------------------------------------------------------

def embed_texts(texts: list[str], model: str = "text-embedding-3-small") -> list[list[float]]:
    """Embed a batch of texts using OpenAI embeddings."""
    from dotenv import load_dotenv
    load_dotenv()
    client = OpenAI()
    all_embeddings: list[list[float]] = []

    batch_size = 100
    for i in range(0, len(texts), batch_size):
        batch = texts[i:i + batch_size]
        response = client.embeddings.create(input=batch, model=model)
        all_embeddings.extend([d.embedding for d in response.data])

    return all_embeddings


# ---------------------------------------------------------------------------
# Full ingestion pipeline
# ---------------------------------------------------------------------------

def ingest_corpus(
    corpus_dir: str | Path,
    index_dir: str | Path | None = None,
) -> tuple[lancedb.DBConnection, nx.DiGraph]:
    """Run the full ingestion pipeline on a corpus directory.

    Returns (lancedb_connection, cross_reference_graph).
    """
    corpus_dir = Path(corpus_dir)
    if index_dir is None:
        index_dir = corpus_dir.parent / "index"
    index_dir = Path(index_dir)
    index_dir.mkdir(parents=True, exist_ok=True)

    files = sorted(
        f for f in os.listdir(corpus_dir)
        if not f.startswith(".") and not f.startswith("~")
    )
    logger.info("Found %d files in %s", len(files), corpus_dir)

    all_chunks: list[Chunk] = []
    all_refs: list[CrossReference] = []
    all_meta: dict[str, DocMetadata] = {}

    for filename in files:
        filepath = corpus_dir / filename
        if filepath.is_dir():
            continue

        meta = parse_filename(filename)
        sections = parse_document(filepath)

        if not sections:
            logger.warning("No content extracted from %s, skipping", filename)
            continue

        full_text = "\n\n".join(text for _, text in sections)

        chunks = chunk_document(meta, sections)
        all_chunks.extend(chunks)

        refs = extract_cross_references(meta, full_text)
        all_refs.extend(refs)

        key = f"{meta.doc_id}_{meta.revision}"
        all_meta[key] = meta

    logger.info("Extracted %d chunks from %d documents", len(all_chunks), len(all_meta))
    logger.info("Found %d cross-references", len(all_refs))

    if not all_chunks:
        raise ValueError("No chunks extracted from corpus")

    # Embed all chunks
    logger.info("Embedding %d chunks...", len(all_chunks))
    texts = [c.text for c in all_chunks]
    vectors = embed_texts(texts)

    # Build LanceDB table
    db = lancedb.connect(str(index_dir / "lancedb"))

    records = []
    for chunk, vector in zip(all_chunks, vectors):
        records.append({
            "text": chunk.text,
            "doc_id": chunk.doc_id,
            "doc_type": chunk.doc_type,
            "title": chunk.title,
            "revision": chunk.revision,
            "status": chunk.status,
            "section": chunk.section,
            "chunk_index": chunk.chunk_index,
            "filename": chunk.filename,
            "vector": vector,
        })

    if "chunks" in db.table_names():
        db.drop_table("chunks")

    table = db.create_table("chunks", data=records)
    table.create_fts_index("text", replace=True)
    logger.info("Created LanceDB table with %d rows and FTS index", len(records))

    # Build cross-reference graph
    graph = nx.DiGraph()

    for key, meta in all_meta.items():
        graph.add_node(meta.doc_id, **{
            "doc_type": meta.doc_type,
            "title": meta.title,
            "revision": meta.revision,
            "status": meta.status,
            "filename": meta.filename,
        })

    for ref in all_refs:
        if not graph.has_node(ref.target_doc_id):
            graph.add_node(ref.target_doc_id)
        graph.add_edge(
            ref.source_doc_id,
            ref.target_doc_id,
            context=ref.context,
            source_file=ref.source_file,
        )

    graph_path = index_dir / "cross_references.pkl"
    with open(graph_path, "wb") as f:
        pickle.dump(graph, f)
    logger.info("Saved cross-reference graph: %d nodes, %d edges", graph.number_of_nodes(), graph.number_of_edges())

    return db, graph


def load_index(
    index_dir: str | Path,
) -> tuple[lancedb.DBConnection, nx.DiGraph]:
    """Load a previously built index from disk."""
    index_dir = Path(index_dir)

    db = lancedb.connect(str(index_dir / "lancedb"))
    if "chunks" not in db.table_names():
        raise FileNotFoundError("LanceDB chunks table not found in index")

    graph_path = index_dir / "cross_references.pkl"
    if not graph_path.exists():
        raise FileNotFoundError("Cross-reference graph not found in index")

    with open(graph_path, "rb") as f:
        graph = pickle.load(f)

    return db, graph


def load_or_build_index(
    corpus_dir: str | Path,
    index_dir: str | Path | None = None,
    force: bool = False,
) -> tuple[lancedb.DBConnection, nx.DiGraph]:
    """Load existing index or build a new one."""
    corpus_dir = Path(corpus_dir)
    if index_dir is None:
        index_dir = corpus_dir.parent / "index"
    index_dir = Path(index_dir)

    if not force and (index_dir / "lancedb").exists() and (index_dir / "cross_references.pkl").exists():
        logger.info("Loading existing index from %s", index_dir)
        return load_index(index_dir)

    logger.info("Building new index from %s", corpus_dir)
    return ingest_corpus(corpus_dir, index_dir)
